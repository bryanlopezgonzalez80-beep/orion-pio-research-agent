from __future__ import annotations

from io import BytesIO
from html import escape
from textwrap import shorten


def _paper_text(p: dict) -> str:
    return (
        f"{p.get('title','')}\n"
        f"Autores: {p.get('authors','')}\n"
        f"Fecha: {p.get('published_date','')}\n"
        f"Fuente: {p.get('journal') or p.get('source','')}\n"
        f"DOI: {p.get('doi','')}\n"
        f"URL: {p.get('url','')}\n\n"
        f"Resumen: {p.get('summary') or p.get('abstract','')}\n\n"
        f"Por qué importa: {p.get('why_it_matters','')}\n\n"
        f"Aplicaciones: {p.get('applications','')}\n\n"
        f"Limitaciones: {p.get('limitations','')}\n\n"
        f"APA: {p.get('apa_citation','')}\n"
    )


def export_docx(papers: list[dict], title: str = "PIO Intelligence Report") -> bytes:
    from docx import Document
    doc = Document()
    doc.add_heading(title, 0)
    doc.add_paragraph("Reporte generado por PIO Intelligence Hub.")
    for i, p in enumerate(papers, 1):
        doc.add_heading(f"{i}. {p.get('title','Sin título')}", level=1)
        doc.add_paragraph(f"Autores: {p.get('authors','')}")
        doc.add_paragraph(f"Fuente: {p.get('journal') or p.get('source','')} | Fecha: {p.get('published_date','')}")
        for heading, field in [
            ("Resumen", "summary"), ("Por qué importa", "why_it_matters"), ("Aplicaciones", "applications"),
            ("Limitaciones", "limitations"), ("Referencia APA", "apa_citation")
        ]:
            doc.add_heading(heading, level=2)
            value = p.get(field) or (p.get("abstract") if field == "summary" else "") or "—"
            doc.add_paragraph(value)
        if p.get("url"):
            doc.add_paragraph(f"Fuente: {p['url']}")
    bio = BytesIO(); doc.save(bio); return bio.getvalue()


def export_pdf(papers: list[dict], title: str = "PIO Intelligence Report") -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from reportlab.lib.units import inch

    bio = BytesIO()
    doc = SimpleDocTemplate(bio, pagesize=letter, rightMargin=0.7*inch, leftMargin=0.7*inch, topMargin=0.7*inch, bottomMargin=0.7*inch)
    styles = getSampleStyleSheet()
    body = ParagraphStyle("Body2", parent=styles["BodyText"], leading=14, spaceAfter=8, alignment=TA_LEFT)
    story = [Paragraph(escape(title), styles["Title"]), Spacer(1, 12), Paragraph("Reporte generado por PIO Intelligence Hub.", body), Spacer(1, 12)]
    for i, p in enumerate(papers, 1):
        story.append(Paragraph(escape(f"{i}. {p.get('title') or 'Sin título'}"), styles["Heading2"]))
        meta = f"{p.get('authors','')} | {p.get('journal') or p.get('source','')} | {p.get('published_date','')}"
        story.append(Paragraph(escape(meta), body))
        for label, value in [
            ("Resumen", p.get("summary") or p.get("abstract") or "—"),
            ("Por qué importa", p.get("why_it_matters") or "—"),
            ("Aplicaciones", p.get("applications") or "—"),
            ("Limitaciones", p.get("limitations") or "—"),
            ("APA", p.get("apa_citation") or "—"),
        ]:
            story.append(Paragraph(f"<b>{label}</b>", body))
            story.append(Paragraph(escape(str(value)).replace("\n", "<br/>"), body))
        if i < len(papers): story.append(Spacer(1, 16))
    doc.build(story)
    return bio.getvalue()


def export_pptx(papers: list[dict], title: str = "PIO Intelligence Brief") -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = "Síntesis de evidencia para consultoría, capacitación y docencia"
    for p in papers[:15]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = shorten(p.get("title", "Sin título"), width=95, placeholder="…")
        tf = slide.placeholders[1].text_frame
        tf.clear()
        items = [
            f"Fuente: {p.get('journal') or p.get('source','')} ({p.get('published_date','')[:4]})",
            f"Qué aporta: {shorten(p.get('summary') or p.get('abstract',''), width=320, placeholder='…')}",
            f"Por qué importa: {shorten(p.get('why_it_matters',''), width=260, placeholder='…')}",
            f"Aplicación: {shorten((p.get('applications') or '').replace(chr(10),' '), width=260, placeholder='…')}",
        ]
        for idx, item in enumerate(items):
            para = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            para.text = item
            para.font.size = Pt(18)
    bio = BytesIO(); prs.save(bio); return bio.getvalue()